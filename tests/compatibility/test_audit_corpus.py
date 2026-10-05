"""Regression checks for the independent audit's comparison boundaries."""

import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Font

from audit_corpus import package_difference, spreadsheet_observations


class AuditTests(unittest.TestCase):
    def test_package_parts_ignore_directory_entries_but_detect_content_loss(self):
        with tempfile.TemporaryDirectory() as directory:
            before, after = (Path(directory) / name for name in ("before.zip", "after.zip"))
            with ZipFile(before, "w") as archive:
                archive.writestr("word/", "")
                archive.writestr("word/document.xml", "original")
            with ZipFile(after, "w") as archive:
                archive.writestr("word/document.xml", "original")
            comparison = package_difference(before, after)
            self.assertTrue(comparison["allPartsByteIdentical"])
            self.assertFalse(comparison["zipDirectoryEntriesEqual"])
            with ZipFile(after, "w") as archive:
                archive.writestr("word/document.xml", "changed")
            comparison = package_difference(before, after)
            self.assertFalse(comparison["allPartsByteIdentical"])
            self.assertEqual(comparison["changedParts"], ["word/document.xml"])

    def test_inherited_bold_and_merged_placeholders_are_not_false_failures(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "sample.xlsx"
            book = Workbook()
            book.active["A1"] = "Title"
            book.active["A1"].font = Font(name="Calibri", size=11, bold=True)
            book.active.merge_cells("A1:B1")
            book.save(source)
            direct = {"italic": "0", "strike": "0", "underline": "0", "wrap": "0",
                      "shrinkToFit": "0", "textRotation": "0", "valign": "bottom",
                      "align": "general", "font": "Calibri", "size": "11"}
            cell = {"address": "A1", "directFormat": direct, "format": {"bold": "1"}}
            covered = {"address": "B1", "directFormat": {"font": "Arial", "size": "32"}}
            report = {"sheets": [{"name": "Sheet", "tables": 0, "cells": [cell, covered]}]}
            comparison = spreadsheet_observations(source, report)["cellComparison"]
            self.assertEqual(comparison["differences"], [])
            self.assertEqual(comparison["coveredMergedCellsNotCompared"], 1)
            self.assertEqual(comparison["fieldsChecked"], 11)
            cell["format"]["bold"] = "0"
            comparison = spreadsheet_observations(source, report)["cellComparison"]
            self.assertEqual([d["field"] for d in comparison["differences"]], ["bold"])
            report["sheets"][0]["tables"] = 1
            comparison = spreadsheet_observations(source, report)["cellComparison"]
            self.assertEqual(comparison["inheritedTableBoldFieldsNotCompared"], 1)
            self.assertEqual(comparison["fieldsChecked"], 10)

    def test_strict_uses_xml_sheet_names_without_claiming_style_checks(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "strict.xlsx"
            with ZipFile(source, "w") as archive:
                archive.writestr("xl/workbook.xml", '<workbook xmlns="http://purl.oclc.org/ooxml/'
                                 'spreadsheetml/main"><sheets><sheet name="Strict"/>'
                                 '</sheets></workbook>')
            result = spreadsheet_observations(source, {"sheets": [{"name": "Strict"}]})
            self.assertTrue(result["sheetNamesMatch"])
            self.assertEqual(result["cellComparison"]["status"], "not_checked")


if __name__ == "__main__":
    unittest.main()
