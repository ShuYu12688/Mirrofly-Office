"""Independent package and selected cell-property audit; never a visual compatibility score."""

import argparse
import hashlib
import json
from pathlib import Path
import warnings
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from pptx import Presentation


def xml(data):
    return etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))


def package_difference(source, destination):
    with ZipFile(source) as before, ZipFile(destination) as after:
        old = {p.filename for p in before.infolist() if not p.is_dir()}
        new = {p.filename for p in after.infolist() if not p.is_dir()}
        changed = [p for p in sorted(old & new) if before.read(p) != after.read(p)]
        return {"partNamesEqual": old == new, "changedParts": changed,
                "missingParts": sorted(old - new), "addedParts": sorted(new - old),
                "allPartsByteIdentical": old == new and not changed,
                "zipDirectoryEntriesEqual":
                    {p.filename for p in before.infolist() if p.is_dir()} ==
                    {p.filename for p in after.infolist() if p.is_dir()}}


def word_observations(source, report):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    model = "\n".join(report.get("paragraphText", []))
    separate = []
    with ZipFile(source) as archive:
        body = xml(archive.read("word/document.xml"))
        for name in archive.namelist():
            if (name.startswith("word/") and name.endswith(".xml") and
                    any(k in name for k in ("header", "footer", "footnotes", "endnotes"))):
                text = xml(archive.read(name)).xpath("//w:t/text()", namespaces=ns)
                separate.append({"part": name,
                                 "nonemptyTextNodes": sum(bool(t.strip()) for t in text),
                                 "textNodesAbsentFromBodyModel":
                                     sum(bool(t.strip()) and t not in model for t in text)})
    return {"separateStoryObservations": separate,
            "textBoxParagraphs": len(body.xpath("//w:txbxContent//w:p", namespaces=ns)),
            "explicitDocumentTabStops": len(body.xpath("//w:tabs/w:tab", namespaces=ns)),
            "probePassed": report.get("success"),
            "editorUnchangedPartsIdentical": report.get("editorUnchangedPartsIdentical"),
            "editedSaveSuccess": report.get("editedSaveSuccess"),
            "editedSaveError": report.get("editedSaveError")}


def spreadsheet_observations(source, report):
    with ZipFile(source) as archive:
        workbook = xml(archive.read("xl/workbook.xml"))
        names = workbook.xpath('//*[local-name()="sheets"]/*[local-name()="sheet"]/@name')
        strict = etree.QName(workbook).namespace == "http://purl.oclc.org/ooxml/spreadsheetml/main"
        rule_types = sorted({node.get("type", "") for name in archive.namelist()
                             if name.startswith("xl/worksheets/") and name.endswith(".xml")
                             for node in xml(archive.read(name)).iter()
                             if isinstance(node.tag, str) and etree.QName(node).localname == "cfRule"})
    result = {"independentXmlSheets": names,
              "sheetNamesMatch": names == [s["name"] for s in report.get("sheets", [])],
              "conditionalRuleTypes": rule_types, "probePassed": report.get("reopened"),
              "observations": [{k: s.get(k) for k in
                                ("name", "storedCells", "truncated", "unsupportedFormulas",
                                 "conditions", "conditionsSupported", "tables", "tablesSupported")}
                               for s in report.get("sheets", [])]}
    # openpyxl's workbook reader does not support Strict sheet relationships; use raw XML above.
    if strict:
        result["cellComparison"] = {"status": "not_checked", "reason": "Strict unsupported by openpyxl"}
        return result
    checked = 0
    skipped_merged = 0
    skipped_table_bold = 0
    differences = []
    with warnings.catch_warnings(record=True) as recorded:
        warnings.simplefilter("always")
        book = load_workbook(source, data_only=False)
        for sheet in report.get("sheets", []):
            if sheet["name"] not in book.sheetnames:
                continue
            for observation in sheet["cells"]:
                actual = observation.get("directFormat")
                if actual is None:
                    continue
                cell = book[sheet["name"]][observation["address"]]
                # openpyxl replaces covered cells with placeholders and discards their stored styles.
                if isinstance(cell, MergedCell):
                    skipped_merged += 1
                    continue
                expected = {"bold": str(int(bool(cell.font.b))), "italic": str(int(bool(cell.font.i))),
                            "strike": str(int(bool(cell.font.strike))),
                            "underline": str(int(bool(cell.font.u))),
                            "wrap": str(int(bool(cell.alignment.wrapText))),
                            "shrinkToFit": str(int(bool(cell.alignment.shrinkToFit))),
                            "textRotation": str(cell.alignment.textRotation or 0),
                            "valign": cell.alignment.vertical or "bottom",
                            "align": cell.alignment.horizontal or "general"}
                if cell.font.name is not None:
                    expected["font"] = cell.font.name
                if cell.font.sz is not None:
                    expected["size"] = float(cell.font.sz)
                for key, value in expected.items():
                    observed = actual.get(key)
                    if key == "bold" and observed is None:
                        # The transferable patch omits inherited bold. Effective format is comparable
                        # only without a table layer; do not manufacture a zero default for font 0.
                        if sheet.get("tables") == 0:
                            observed = observation.get("format", {}).get(key)
                        else:
                            skipped_table_bold += 1
                            continue
                    if key == "size":
                        observed = float(observed) if observed is not None else None
                    checked += 1
                    if observed != value:
                        differences.append({"sheet": sheet["name"], "cell": observation["address"],
                                            "field": key, "expected": value, "actual": observed})
        result["independentReaderWarnings"] = sorted({str(w.message) for w in recorded})
    result["cellComparison"] = {"status": "checked", "fieldsChecked": checked,
                                "coveredMergedCellsNotCompared": skipped_merged,
                                "inheritedTableBoldFieldsNotCompared": skipped_table_bold,
                                "differences": differences,
                                "scope": "Listed direct font/alignment fields only; not number display, "
                                         "theme colors, table/conditional layers or rendered pixels."}
    return result


def audit(fixtures, results, spreadsheet_results=None):
    manifest = json.loads((fixtures / "manifest.json").read_text(encoding="utf-8"))
    observations = []
    for item in manifest["files"]:
        source = fixtures / item["path"]
        group = source.parent.name
        result_root = spreadsheet_results if group == "spreadsheet" and spreadsheet_results else results
        directory = result_root / group / source.stem
        row = {"file": item["path"], "sourceUnchanged":
               hashlib.sha256(source.read_bytes()).hexdigest() == item["sha256"]}
        report_path = directory / "report.json"
        if not report_path.exists():
            row["missingReport"] = str(report_path)
            observations.append(row)
            continue
        report = json.loads(report_path.read_text(encoding="utf-8"))
        roundtrips = list(directory.rglob("roundtrip" + source.suffix))
        row["roundtripPresent"] = len(roundtrips) == 1
        if len(roundtrips) == 1:
            row.update(package_difference(source, roundtrips[0]))
        if group == "document":
            row.update(word_observations(source, report))
        elif group == "spreadsheet":
            row.update(spreadsheet_observations(source, report))
        else:
            presentation = Presentation(source)
            row.update({"independentSlideCount": len(presentation.slides),
                        "coreSlideCount": len(report[0].get("slides", [])),
                        "probePassed": report[0].get("saveError") == 0,
                        "loadError": report[0].get("message", ""),
                        "warnings": report[0].get("warnings", []),
                        "slideWarnings": [s.get("warnings", []) for s in report[0].get("slides", [])
                                          if s.get("warnings")]})
        observations.append(row)
    return {"repositoryRevision": manifest["revision"],
            "scope": "Feature-selected inputs, not a population fidelity or functional coverage score.",
            "results": observations}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fixtures", required=True, type=Path)
    parser.add_argument("--results", required=True, type=Path)
    parser.add_argument("--spreadsheet-results", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    result = audit(args.fixtures, args.results, args.spreadsheet_results)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    for group in ("document", "spreadsheet", "slideshow"):
        rows = [r for r in result["results"] if r["file"].startswith(group + "/")]
        print(group, "samples", len(rows), "identical file parts",
              sum(bool(r.get("allPartsByteIdentical")) for r in rows), "probe passes",
              sum(bool(r.get("probePassed")) for r in rows))
