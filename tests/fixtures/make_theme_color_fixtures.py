"""Independent DOCX/XLSX fixtures for theme-aware parsing and public AI reads."""
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Color, Font, PatternFill, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.table import TableStyle, TableStyleElement
from openpyxl.worksheet.table import Table, TableStyleInfo

root = Path(__file__).parent
word = Document()
p = word.add_paragraph()
r = p.add_run("主题颜色示例 Theme color")
c = OxmlElement("w:color")
c.set(qn("w:val"), "000000")
c.set(qn("w:themeColor"), "accent1")
c.set(qn("w:themeTint"), "99")
r._r.get_or_add_rPr().append(c)
p = word.add_paragraph("保留主题底色和边框，仅修改标题字号。")
shd = OxmlElement("w:shd")
for k,v in {"val":"clear","fill":"000000","themeFill":"accent1","themeFillTint":"99"}.items():
    shd.set(qn("w:"+k),v)
p._p.get_or_add_pPr().append(shd)
word.save(root / "word-theme-colors.docx")
book = Workbook()
book.properties.created = datetime(2026,9,29)
book.properties.modified = datetime(2026,9,29)
s = book.active
s.title = "主题颜色"
s.append(["标题", "数值", "说明"])
s.append(["保留主题", 12, "只修改A1字号与A2数值"])
s["A1"].font = Font(name="Calibri",color=Color(theme=4,tint=.4),size=14)
s["A1"].fill = PatternFill("solid",fgColor=Color(theme=4,tint=.8))
s["A1"].border = Border(bottom=Side(style="thin",color=Color(indexed=10)))
s["B1"].font = Font(color=Color(theme=1))
s["B1"].fill = PatternFill("solid",fgColor=Color(theme=0))
s["A2"].fill = PatternFill("solid",fgColor=Color(indexed=10))
s.conditional_formatting.add("B2",CellIsRule(operator="greaterThan",formula=[10],fill=PatternFill("solid",fgColor=Color(theme=4,tint=.4))))
s.column_dimensions["A"].width=24
s.column_dimensions["C"].width=32
s["D1"], s["E1"] = "项目", "数量"
s["D2"], s["E2"] = "样例", 3
style_id = book._differential_styles.add(DifferentialStyle(fill=PatternFill("solid",fgColor=Color(theme=4,tint=.4))))
book._table_styles.tableStyle.append(TableStyle(name="ThemeStyle",pivot=False,table=True,tableStyleElement=[TableStyleElement(type="wholeTable",dxfId=style_id)]))
table = Table(displayName="ThemeTable",ref="D1:E2")
table.tableStyleInfo = TableStyleInfo(name="ThemeStyle",showRowStripes=False)
s.add_table(table)
book.save(root / "spreadsheet-theme-colors.xlsx")
