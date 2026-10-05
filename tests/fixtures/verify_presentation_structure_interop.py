"""Independent validation of corpus-probe structure edits (never writes input decks)."""
import argparse
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation


NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def xml(data):
    return etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True,
                                                 remove_blank_text=True))


def canonical(node):
    return etree.tostring(node, method='c14n')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('probe', type=Path)
    args = parser.parse_args()
    report = json.loads((args.probe / 'report.json').read_text(encoding='utf-8'))[0]
    edited = next(args.probe.rglob('structure-edit.pptx'))
    roundtrip = edited.with_name('roundtrip.pptx')
    with ZipFile(args.source) as original, ZipFile(roundtrip) as copy, ZipFile(edited) as changed:
        assert set(original.namelist()) == set(copy.namelist()) == set(changed.namelist())
        differences = set()
        for name in original.namelist():
            before = original.read(name)
            assert before == copy.read(name), f'unedited part changed: {name}'
            if before != changed.read(name):
                differences.add(name)
        assert differences == {item['part'] for item in report['structureChanges']}
        for item in report['structureChanges']:
            before, after = xml(original.read(item['part'])), xml(changed.read(item['part']))
            if not item['cell']:
                query = './/p:cNvPr[@id=$id]/../..'
                old = before.xpath(query, namespaces=NS, id=item['sourceId'])[0]
                new = after.xpath(query, namespaces=NS, id=item['sourceId'])[0]
                old_xfrm, new_xfrm = old.find('p:spPr/a:xfrm', NS), new.find('p:spPr/a:xfrm', NS)
                assert old_xfrm is not None and new_xfrm is not None
                assert canonical(old_xfrm) != canonical(new_xfrm)
                old_xfrm.getparent().remove(old_xfrm)
                new_xfrm.getparent().remove(new_xfrm)
                assert canonical(before) == canonical(after), 'non-target group properties changed'
            else:
                old_table, new_table = before.find('.//a:tbl', NS), after.find('.//a:tbl', NS)
                assert canonical(old_table.find('a:tblGrid', NS)) == canonical(new_table.find('a:tblGrid', NS))
                for old_row, new_row in zip(old_table.findall('a:tr', NS), new_table.findall('a:tr', NS)):
                    assert old_row.attrib == new_row.attrib
                    for old_cell, new_cell in zip(old_row.findall('a:tc', NS), new_row.findall('a:tc', NS)):
                        assert old_cell.attrib == new_cell.attrib
                        assert canonical(old_cell.find('a:tcPr', NS)) == canonical(new_cell.find('a:tcPr', NS))
                assert new_table.xpath('.//a:tr[1]/a:tc[1]//a:rPr/a:solidFill/a:srgbClr/@val', namespaces=NS)
                assert set(new_table.xpath('.//a:tr[1]/a:tc[1]//a:rPr/a:solidFill/a:srgbClr/@val', namespaces=NS)) == {'13579B'}
        print(f'ZIP: {len(original.namelist())} untouched-copy parts byte identical; edits limited to {sorted(differences)}.')
    loaded = Presentation(edited)
    source = Presentation(args.source)
    assert len(loaded.slides) == len(source.slides)
    print(f'python-pptx reopened {len(loaded.slides)} slides; group ancestors, table grid, row heights and cell properties preserved.')


if __name__ == '__main__':
    main()
