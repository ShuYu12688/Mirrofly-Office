"""Independent SpreadsheetML text-direction fixture; no app-generated styles."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

book = Workbook()
sheet = book.active
sheet.title = "文字方向"
for row, angle in enumerate([0, 45, 90, 135, 180, 255, 30, 150, 45, 90], 1):
    cell = sheet.cell(row, 1, "竖排Abc" if angle == 255 else f"方向{row} Mixed 123")
    cell.font = Font(name="Calibri", size=20, color="234567")
    cell.fill = PatternFill("solid", fgColor="E8F0FA")
    cell.alignment = Alignment(horizontal="center", vertical="center", textRotation=angle,
                               wrapText=row == 9, shrinkToFit=row == 10)
    sheet.row_dimensions[row].height = 120
sheet["A9"] = "自动换行 Mixed words that wrap across several lines"
sheet["A10"] = "旋转缩小 Mixed long line ABCDEFGHIJKLMNOPQRSTUVWXYZ"
sheet.column_dimensions["A"].width = 30
sheet.column_dimensions["B"].width = 20
sheet.merge_cells("A8:B8")
book.save(Path(__file__).with_name("spreadsheet-rotation.xlsx"))
