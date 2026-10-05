from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


def verify():
    project = Path(__file__).resolve().parents[2]
    original = Path(__file__).with_name("spreadsheet-openpyxl.xlsx")
    output = project / "build/spreadsheet-interop-result.xlsx"
    book = load_workbook(output, data_only=False)
    assert book.sheetnames == ["进度", "明细"]
    sheet = book["进度"]
    assert sheet["A2"].value == "Mirrorfly 编辑"
    assert sheet["B2"].value == 88.5
    assert sheet["C2"].value == "001234"
    assert sheet["A3"].value == "=SUM(B2,1)"
    assert sheet["A2"].font.bold
    assert sheet["A2"].fill.fgColor.rgb == "FFEAF2EC"
    assert sheet["A2"].alignment.horizontal == "center"
    assert sheet["B2"].border.bottom.style == "thin"
    assert sheet.column_dimensions["A"].width == 26
    assert sheet.row_dimensions[2].height == 40
    assert sheet["B2"].font.bold
    assert sheet["B2"].number_format == "0.00"
    assert sheet["B2"].comment.text == "独立生成的保留部件检查"
    assert "D5:E5" in sheet.merged_cells
    assert book["明细"]["A2"].value.year == 2026
    assert book["明细"]["B2"].value is True
    assert book.calculation.fullCalcOnLoad and book.calculation.forceFullCalc
    assert load_workbook(output, data_only=True)["进度"]["A3"].value == 89.5
    with ZipFile(original) as before, ZipFile(output) as after:
        untouched = [name for name in before.namelist()
                     if name not in ("xl/workbook.xml", "xl/styles.xml") and not name.startswith("xl/worksheets/sheet")]
        for name in untouched:
            assert before.read(name) == after.read(name), name
    extended = load_workbook(project / "build/spreadsheet-start-tools.xlsx")
    e = extended.active
    assert e["A2"].font.underline == "single" and e["A2"].font.strike
    assert e["A2"].number_format == '"¥"#,##0.000'
    assert e["A2"].alignment.vertical == "top" and e["A2"].alignment.indent == 2
    assert e["A2"].alignment.horizontal == "distributed"
    assert e["B3"].alignment.horizontal == "justify" and e["B3"].alignment.wrap_text
    assert "A5:C5" in e.merged_cells and e.freeze_panes == "A2"
    assert e.auto_filter.ref == "A1:B3" and e.auto_filter.filterColumn[0].filters.filter == ["5"]
    assert len(e.auto_filter.filterColumn) == 2
    assert e.auto_filter.filterColumn[1].colId == 1
    assert e.auto_filter.filterColumn[1].customFilters.customFilter[0].val == "*~*~?*"
    assert e["A2"].border.right.style == "double" and e["A2"].border.right.color.rgb == "FF123456"
    assert e["A2"].border.bottom.style == "dashed" and e["A2"].border.bottom.color.rgb == "FFCC2211"
    assert e["A2"].border.left.style is None and e["A2"].border.top.style is None
    assert e.row_dimensions[3].hidden and not e.row_dimensions[2].hidden
    assert e.column_dimensions["E"].hidden
    rule = list(e.conditional_formatting._cf_rules.values())[0][0]
    assert rule.operator == "greaterThan" and rule.formula == ["6"] and rule.dxf.fill.fgColor.rgb == "FFAABBCC"
    tables = load_workbook(project / "build/spreadsheet-tables.xlsx")
    t = tables.active.tables["ProgressTable"]
    assert t.ref == "A1:B3" and t.displayName == "ProgressTable"
    assert [column.name for column in t.tableColumns] == ["数量", "名称"]
    assert t.tableStyleInfo.showRowStripes and t.tableStyleInfo.showFirstColumn
    assert not tables.active.auto_filter.ref
    assert t.autoFilter.ref == "A1:B3" and len(t.autoFilter.filterColumn) == 2
    assert t.autoFilter.filterColumn[1].customFilters.customFilter[0].val == "*~*~?*"
    style = next(style for style in tables._table_styles.tableStyle if style.name == t.tableStyleInfo.name)
    elements = {element.type: tables._differential_styles.styles[element.dxfId]
                for element in style.tableStyleElement}
    assert elements["headerRow"].fill.fgColor.rgb == "FF35644C"
    assert elements["headerRow"].font.bold and elements["headerRow"].font.color.rgb == "FFFFFFFF"
    assert elements["firstRowStripe"].fill.fgColor.rgb == "FFEDF4EF"
    assert elements["firstColumn"].font.bold
    assert tables.active["A2"].border.right.style == "double"
    assert tables.active.row_dimensions[3].hidden
    assert not tables.active["A1"].has_style  # Table formatting remains a separate style layer.
    structure_path = project / "build/spreadsheet-structure.xlsx"
    structure = load_workbook(structure_path)
    assert structure.sheetnames == ["副本", "汇总", "数据"]
    data = structure["数据"]
    assert data.sheet_state == "hidden" and structure["副本"].sheet_state == "visible"
    assert data["A2"].value == 5 and data["A4"].value == 9 and data["A3"].value is None
    assert data["B1"].value == "列1" and data["C1"].value == "名称"
    assert "A6:D6" in data.merged_cells and data.freeze_panes == "A2"
    assert data.column_dimensions["F"].hidden and data.row_dimensions[4].hidden
    assert data["A2"].border.right.style == "double"
    original_table = data.tables["ProgressTable"]
    copied_table = next(iter(structure["副本"].tables.values()))
    assert original_table.ref == copied_table.ref == "A1:C4"
    assert original_table.id != copied_table.id and original_table.name != copied_table.name
    assert [c.name for c in original_table.tableColumns] == ["数量", "列1", "名称"]
    assert original_table.autoFilter.filterColumn[1].colId == 2
    assert structure["汇总"]["A1"].value == "=SUM('数据'!A2:A4)"
    assert load_workbook(structure_path, data_only=True)["汇总"]["A1"].value == 14
    assert structure.defined_names["Amount"].attr_text == "'数据'!$A$2:$A$4"
    assert data.defined_names["LocalAmount"].attr_text == "'数据'!$A$2"
    assert data.defined_names["LocalAmount"].localSheetId == 2
    assert structure["副本"].defined_names["LocalAmount"].attr_text == "'副本'!$A$2"
    assert structure["副本"].defined_names["LocalAmount"].localSheetId == 0
    moved = load_workbook(project / "build/spreadsheet-moved.xlsx")
    assert not moved["副本"].tables and not moved["副本"].merged_cells
    target = moved["汇总"]
    moved_table = next(iter(target.tables.values()))
    assert moved_table.ref == "E3:G6" and moved_table.autoFilter.ref == "E3:G6"
    assert target["E4"].value == 5 and target["E6"].value == 9
    assert target["E4"].border.right.style == "double" and "E8:H8" in target.merged_cells
    assert not target.row_dimensions[4].hidden and target.row_dimensions[6].hidden
    assert not moved["副本"].row_dimensions[4].hidden
    assert moved["副本"].defined_names["LocalAmount"].attr_text == "'汇总'!$E$4"
    assert list(target.conditional_formatting)[0].sqref == "E4:E6"
    assert moved["数据"].tables["ProgressTable"].ref == "A1:C4"
    print("openpyxl 3.1.5 independently verified values, styles, comments, dates, formulas and untouched parts.")


if __name__ == "__main__":
    verify()
