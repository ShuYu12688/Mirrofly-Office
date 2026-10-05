"""Independent Word table theme and SpreadsheetML shrink fixtures."""
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

root = Path(__file__).parent
word = Document()
word.add_paragraph("表格主题颜色核验")
table = word.add_table(rows=2, cols=2)
for cell, value in zip([c for row in table.rows for c in row.cells],
                       ["主题单元格", "保持内容", "第二行", "123"]):
    cell.text = value
shade = OxmlElement("w:shd")
for key, value in {"val": "clear", "fill": "000000", "themeFill": "accent1",
                   "themeFillTint": "99"}.items():
    shade.set(qn("w:" + key), value)
table.cell(0, 0)._tc.get_or_add_tcPr().append(shade)
borders = OxmlElement("w:tblBorders")
for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
    border = OxmlElement("w:" + edge)
    for key, value in {"val": "single", "sz": "8", "color": "000000",
                       "themeColor": "accent1"}.items():
        border.set(qn("w:" + key), value)
    borders.append(border)
table._tbl.tblPr.append(borders)
word.save(root / "word-table-theme.docx")

book = Workbook()
sheet = book.active
sheet.title = "文字布局"
for row, text in enumerate(["很长的单行标题 Mixed text 123456789", "待启用缩小的长文本 ABCDEFGHIJKLMN",
                            "换行优先 Multi line sample", "第一行内容\n第二行内容", "合并区域的长标题 ABCDEFGHIJKLMNOP"], 1):
    sheet.cell(row, 1, text)
    sheet.cell(row, 1).font = Font(name="Calibri", size=20)
    sheet.cell(row, 1).alignment = Alignment(horizontal="left", vertical="center",
                                            shrinkToFit=row != 2, wrapText=row == 3)
    sheet.row_dimensions[row].height = 40
sheet.column_dimensions["A"].width = 12
sheet.column_dimensions["B"].width = 12
sheet.merge_cells("A5:B5")
book.save(root / "spreadsheet-shrink.xlsx")
