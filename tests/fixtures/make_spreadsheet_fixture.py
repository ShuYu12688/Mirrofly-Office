from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill


def make_fixture():
    book = Workbook()
    book.properties.creator = "Mirrorfly developer tests"
    book.properties.created = datetime(2026, 9, 15)
    book.properties.modified = datetime(2026, 9, 15)
    progress = book.active
    progress.title = "进度"
    progress.append(["项目", "进度", "编号"])
    progress.append(["初稿", 50.25, "001234"])
    progress["A3"] = "=SUM(B2,1)"
    progress["B2"].number_format = "0.00"
    progress["B2"].font = Font(name="Arial", bold=True, color="315D48")
    progress["B2"].fill = PatternFill("solid", fgColor="E7EFE5")
    progress["B2"].comment = Comment("独立生成的保留部件检查", "Mirrorfly")
    progress["D5"] = "合并区域"
    progress.merge_cells("D5:E5")
    progress.column_dimensions["A"].width = 26
    details = book.create_sheet("明细")
    details.append(["日期", "状态"])
    details.append([datetime(2026, 9, 15), True])
    details["A2"].number_format = "yyyy-mm-dd"
    book.save(Path(__file__).with_name("spreadsheet-openpyxl.xlsx"))


if __name__ == "__main__":
    make_fixture()
