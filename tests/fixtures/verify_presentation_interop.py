"""Independent python-pptx and XML checks for the generated text-style fixture."""

from pathlib import Path
import posixpath
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def verify():
    path = Path(__file__).resolve().parents[2] / 'build' / 'presentation-text-style.pptx'
    deck = Presentation(path)
    assert len(deck.slides) == 1
    assert deck.slides[0].shapes[0].has_text_frame
    assert deck.slides[0].shapes[0].text_frame.text
    namespace = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
                 'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with ZipFile(path) as package:
        assert package.testzip() is None
        root = etree.fromstring(package.read('ppt/slides/slide1.xml'), parser)
        run = root.find('.//a:rPr', namespace)
        gradient = run.find('a:gradFill', namespace)
        assert gradient is not None
        stops = gradient.findall('a:gsLst/a:gs', namespace)
        assert [int(stop.get('pos')) for stop in stops] == [0, 100000]
        assert [stop.find('a:srgbClr', namespace).get('val') for stop in stops] == ['FF0000', '0000FF']
        assert int(gradient.find('a:lin', namespace).get('ang')) == 90 * 60000
        effects = run.find('a:effectLst', namespace)
        shadow = effects.find('a:outerShdw', namespace)
        assert shadow.find('a:srgbClr', namespace).get('val') == '111111'
        reflection = effects.find('a:reflection', namespace)
        assert int(reflection.get('stA')) == 55000
        assert int(reflection.get('endPos')) == 45500
        assert int(reflection.get('sy')) == -100000
        run_order = {'ln': 0, 'noFill': 1, 'solidFill': 1, 'gradFill': 1, 'effectLst': 2,
                     'highlight': 3, 'latin': 6, 'ea': 7, 'cs': 8, 'sym': 9}
        positions = [run_order[etree.QName(child).localname] for child in run
                     if etree.QName(child).localname in run_order]
        assert positions == sorted(positions), 'DrawingML run properties must follow schema order'
        style_list = etree.fromstring(package.read('ppt/tableStyles.xml'), parser)
        style = style_list.find('a:tblStyle', namespace)
        assert style is not None and style.get('styleId') == style_list.get('def')
        for region in ('wholeTbl', 'firstRow', 'lastRow', 'firstCol', 'lastCol',
                       'band1H', 'band2H', 'band1V', 'band2V'):
            assert style.find('a:' + region + '/a:tcStyle/a:fill/a:solidFill/a:schemeClr', namespace) is not None
        assert style.find('a:firstRow/a:tcStyle/a:fill/a:solidFill/a:schemeClr', namespace).get('val') == 'accent1'
        assert style.find('a:firstRow/a:tcTxStyle/a:fontRef', namespace).get('idx') == 'minor'
    table_path = path.with_name('presentation-table-structure.pptx')
    structured = Presentation(table_path)
    assert len(structured.slides) == 1
    tables = [shape.table for shape in structured.slides[0].shapes if shape.has_table]
    assert len(tables) == 1 and len(tables[0].rows) == 4 and len(tables[0].columns) == 3
    assert tables[0].cell(0, 0).is_merge_origin and tables[0].cell(0, 0).span_width == 2
    assert tables[0].cell(1, 0).is_merge_origin and tables[0].cell(1, 0).span_width == 2
    with ZipFile(table_path) as package:
        assert package.testzip() is None
        root = etree.fromstring(package.read('ppt/slides/slide1.xml'), parser)
        rows = root.findall('.//a:tbl/a:tr', namespace)
        assert len(rows) == 4
        for row in rows[:2]:
            cells = row.findall('a:tc', namespace)
            assert len(cells) == 3 and cells[0].get('gridSpan') == '2'
            assert cells[1].get('hMerge') == '1'
    transition_path = path.with_name('presentation-transition.pptx')
    transitioned = Presentation(transition_path)
    assert len(transitioned.slides) == 2
    with ZipFile(transition_path) as package:
        assert package.testzip() is None
        first = etree.fromstring(package.read('ppt/slides/slide1.xml'), parser)
        second = etree.fromstring(package.read('ppt/slides/slide2.xml'), parser)
        assert first.find('p:transition', namespace) is None
        transition = second.find('p:transition', namespace)
        assert transition is not None
        assert transition.get('spd') == 'slow' and transition.get('advClick') == '0'
        assert transition.get('advTm') == '2500'
        assert transition.find('p:push', namespace).get('dir') == 'l'
    group_path = path.with_name('presentation-group-layer.pptx')
    grouped = Presentation(group_path)
    assert len(grouped.slides) == 1
    assert grouped.slides[0].shapes[-1].shape_type == MSO_SHAPE_TYPE.GROUP
    assert len(grouped.slides[0].shapes[-1].shapes) == 2
    with ZipFile(group_path) as package:
        assert package.testzip() is None
        root = etree.fromstring(package.read('ppt/slides/slide1.xml'), parser)
        tree = root.find('p:cSld/p:spTree', namespace)
        order = [etree.QName(item).localname for item in tree
                 if etree.QName(item).localname in ('sp', 'pic', 'grpSp', 'graphicFrame', 'cxnSp')]
        assert order == ['sp', 'grpSp']
    theme_path = path.with_name('presentation-theme-state.pptx')
    themed = Presentation(theme_path)
    assert len(themed.slides) == 3

    def related_part(package, owner, kind):
        relations_path = posixpath.join(posixpath.dirname(owner), '_rels',
                                       posixpath.basename(owner) + '.rels')
        relations = etree.fromstring(package.read(relations_path), parser)
        matches = [item for item in relations if item.get('Type', '').endswith('/' + kind)]
        assert len(matches) == 1
        target = matches[0].get('Target')
        assert target and matches[0].get('TargetMode') != 'External'
        return posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))

    with ZipFile(theme_path) as package:
        assert package.testzip() is None
        theme_parts = []
        for slide_number in (1, 2, 3):
            slide_part = f'ppt/slides/slide{slide_number}.xml'
            layout_part = related_part(package, slide_part, 'slideLayout')
            master_part = related_part(package, layout_part, 'slideMaster')
            theme_parts.append(related_part(package, master_part, 'theme'))
        assert len(set(theme_parts)) == 3
        theme = etree.fromstring(package.read(theme_parts[1]), parser)
        assert theme.find('.//a:clrScheme/a:accent1/a:srgbClr', namespace).get('val') == '123456'
        assert theme.find('.//a:fontScheme/a:majorFont/a:latin', namespace).get('typeface') == 'Georgia'
    print('PPTX interop: python-pptx reads text, table, transitions, groups and themes; XML verifies structure.')


if __name__ == '__main__':
    verify()
