"""Read actual AI-edited Unicode DOCX with python-docx and an independent XML reader."""

from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document


root = Path(__file__).resolve().parents[2]
original = root / "tests/fixtures/word-unicode.docx"
edited = root / "build/word-unicode-edited.docx"
expected = ["中文 A🦋B𠀀 末尾e\u0301", "保留段落 Latin"]
document = Document(edited)
assert [p.text for p in document.paragraphs] == expected
assert document.paragraphs[1].runs[0].italic
assert document.sections[0].header.paragraphs[0].text == "Preserved header"
offset = 0
for run in document.paragraphs[0].runs:
    end = offset + len(run.text.encode("utf-16-le")) // 2
    selected = 4 <= offset and end <= 9
    assert not (offset < 4 < end or offset < 9 < end)
    assert run.font.size.pt == (12.5 if selected else 10.5)
    assert bool(run.bold) == selected
    offset = end
assert offset == len(expected[0].encode("utf-16-le")) // 2
with ZipFile(original) as source, ZipFile(edited) as saved:
    assert saved.testzip() is None
    assert set(source.namelist()) == set(saved.namelist())
    assert [name for name in source.namelist() if source.read(name) != saved.read(name)] == ["word/document.xml"]
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    before = ET.fromstring(source.read("word/document.xml")).findall("w:body/w:p", ns)
    after = ET.fromstring(saved.read("word/document.xml")).findall("w:body/w:p", ns)
    assert ET.tostring(before[1]) == ET.tostring(after[1])
print("Word Unicode/half-point size: AI output, UTF-16 range, original parts and independent readback passed.")
