"""Independent mixed Unicode and half-point-size DOCX fixture (python-docx)."""

from pathlib import Path
from docx import Document
from docx.shared import Pt


document = Document()
paragraph = document.add_paragraph()
for text in ("中文 A", "🦋B𠀀", " 末尾e\u0301"):
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.bold = False
document.add_paragraph("保留段落 Latin").runs[0].italic = True
document.sections[0].header.paragraphs[0].text = "Preserved header"
document.save(Path(__file__).with_name("word-unicode.docx"))
