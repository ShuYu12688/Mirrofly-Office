"""Independent python-docx table style fixture; source has no flattened cell fonts."""
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

root = Path(__file__).parent
doc = Document()
doc.add_paragraph("表格样式继承验证：保留结构、主题和原始样式定义。")
base = doc.styles.add_style("Mirrorfly Table Base", WD_STYLE_TYPE.TABLE)
for xml in [
    '<w:pPr><w:spacing w:after="0"/></w:pPr>',
    '<w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:eastAsia="宋体"/>'
    '<w:color w:val="234567"/><w:sz w:val="32"/></w:rPr>',
    '<w:tblPr><w:tblStyleRowBandSize w:val="1"/><w:tblStyleColBandSize w:val="1"/>'
    '<w:tblBorders>' + ''.join(f'<w:{edge} w:val="single" w:sz="8" w:color="4F81BD"/>'
                             for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']) +
    '</w:tblBorders><w:tblCellMar><w:top w:w="80"/><w:left w:w="120"/>'
    '<w:bottom w:w="80"/><w:right w:w="120"/></w:tblCellMar></w:tblPr>',
    '<w:tcPr><w:shd w:fill="FFFFFF"/></w:tcPr>',
    '<w:tblStylePr w:type="firstRow"><w:rPr><w:b/><w:color w:val="FFFFFF"/>'
    '<w:sz w:val="36"/></w:rPr><w:tcPr><w:shd w:fill="000000" w:themeFill="accent1"/>'
    '</w:tcPr></w:tblStylePr>',
    '<w:tblStylePr w:type="band1Horz"><w:tcPr><w:shd w:fill="000000" '
    'w:themeFill="accent1" w:themeFillTint="99"/></w:tcPr></w:tblStylePr>',
    '<w:tblStylePr w:type="band2Horz"><w:tcPr><w:shd w:fill="E9EEF4"/></w:tcPr></w:tblStylePr>',
    '<w:tblStylePr w:type="lastRow"><w:rPr><w:b/></w:rPr>'
    '<w:tcPr><w:shd w:fill="D9EAD3"/></w:tcPr></w:tblStylePr>',
]:
    base.element.append(parse_xml(xml.replace('>', ' ' + nsdecls('w') + '>', 1)))
derived = doc.styles.add_style("Mirrorfly Table Derived", WD_STYLE_TYPE.TABLE)
derived.base_style = base
table = doc.add_table(rows=4, cols=3)
table.style = derived
old = table._tbl.tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblLook')
if old is not None:
    table._tbl.tblPr.remove(old)
table._tbl.tblPr.append(parse_xml('<w:tblLook ' + nsdecls('w') +
    ' w:firstRow="1" w:lastRow="1" w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/>'))
for row, values in zip(table.rows, [
        ['表头甲', '表头乙', '表头丙'], ['数据甲1', '数据乙1', '数据丙1'],
        ['数据甲2', '数据乙2', '数据丙2'], ['合计甲', '合计乙', '合计丙']]):
    for cell, text in zip(row.cells, values):
        cell.text = text
table.cell(1, 1)._tc.get_or_add_tcPr().append(parse_xml(
    '<w:shd ' + nsdecls('w') + ' w:val="clear" w:fill="FFF2CC"/>'))
doc.add_paragraph("尾段保持原样。")
doc.save(root / 'word-table-styles.docx')
