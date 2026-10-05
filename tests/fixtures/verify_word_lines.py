"""Read the actual QML/AI Word bridge output with python-docx and an independent XML parser."""

from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import sys

from docx import Document
from docx.enum.text import WD_UNDERLINE


path = Path(sys.argv[1]) if len(sys.argv) > 1 else (
    Path(__file__).resolve().parents[2] / 'build' / 'word-double-lines.docx')
document = Document(path)
assert len(document.paragraphs) == 1
paragraph = document.paragraphs[0]
assert paragraph.text == 'target keep'
assert [(r.text, r.font.underline, bool(r.font.strike), bool(r.font.double_strike))
        for r in paragraph.runs] == [
    ('target', WD_UNDERLINE.DOUBLE, False, True),
    (' ', False, False, False),
    ('keep', True, True, False),
]
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
attribute = '{' + ns['w'] + '}val'
with ZipFile(path) as archive:
    assert archive.testzip() is None
    root = ET.fromstring(archive.read('word/document.xml'))
    for properties in root.findall('.//w:rPr', ns):
        def enabled(name):
            node = properties.find('w:' + name, ns)
            return node is not None and node.get(attribute, '1') not in ('0', 'false', 'off')
        assert not (enabled('strike') and enabled('dstrike'))
print('Word lines: actual parent UI double + actual local AI single + unchanged space; independent DOCX passed.')
