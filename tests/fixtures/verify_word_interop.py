"""Independently inspect a DOCX written by the C++ workbench regression test."""

from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document


path = Path(__file__).resolve().parents[2] / "build" / "word-interop.docx"
with ZipFile(path) as archive:
    assert archive.testzip() is None
document = Document(path)
assert [paragraph.text for paragraph in document.paragraphs] == ["项目计划", "第一阶段：完成评审。", "第二阶段：交付。"]
run = document.paragraphs[0].runs[0]
assert run.font.size.pt == 24 and run.bold
assert run.font.name == "Arial"
assert run.font.strike and run.font.superscript
assert str(run.font.color.rgb) == "123456"
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
with ZipFile(path) as archive:
    root = ET.fromstring(archive.read("word/document.xml"))
    p = root.find("w:body/w:p", ns)
    w = "{" + ns["w"] + "}"
    assert p.find("w:pPr/w:shd", ns).get(w + "fill").upper() == "CCDDEE"
    assert p.find("w:pPr/w:pBdr/w:bottom", ns).get(w + "color") == "445566"
    assert p.find("w:r/w:rPr/w:shd", ns).get(w + "fill").upper() == "FFEEDD"
    assert p.find("w:r/w:rPr/w:bdr", ns).get(w + "color").upper() == "667788"
    assert p.find("w:r/w:rPr/w:spacing", ns).get(w + "val") == "30"
assert abs(document.sections[0].page_width.mm - 210) < 1
assert abs(document.sections[0].page_height.mm - 297) < 1
formatting = document.paragraphs[0].paragraph_format
assert formatting.left_indent.pt == 36
assert formatting.first_line_indent.pt == -18
assert formatting.space_before.pt == 9 and formatting.space_after.pt == 12

template = path.with_name("word-workflows-interop.docx")
meeting = Document(template)
assert meeting.paragraphs[0].text == "会议纪要"
assert len(meeting.paragraphs) == 10
assert "[待填写]" in meeting.paragraphs[-1].text
namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
attribute = "{" + namespace["w"] + "}"
with ZipFile(template) as archive:
    assert archive.testzip() is None
    numbering = ET.fromstring(archive.read("word/numbering.xml"))
    content = ET.fromstring(archive.read("word/document.xml"))
    numbers = {item.attrib[attribute + "numId"]:
               item.find("w:abstractNumId", namespace).attrib[attribute + "val"]
               for item in numbering.findall("w:num", namespace)}
    abstracts = {item.attrib[attribute + "abstractNumId"]: item
                 for item in numbering.findall("w:abstractNum", namespace)}
    paragraphs = content.findall("w:body/w:p", namespace)
    for index, expected in [(5, "bullet"), (7, "decimal")]:
        properties = paragraphs[index].find("w:pPr/w:numPr", namespace)
        number = properties.find("w:numId", namespace).attrib[attribute + "val"]
        level = properties.find("w:ilvl", namespace).attrib[attribute + "val"]
        entry = abstracts[numbers[number]].find(f"w:lvl[@w:ilvl='{level}']/w:numFmt", namespace)
        assert entry.attrib[attribute + "val"] == expected
print("Independent DOCX read: text, font, A4 geometry, paragraph settings and real template lists passed.")
with ZipFile(path.with_name("word-phonetic-interop.docx")) as archive:
    ruby_root = ET.fromstring(archive.read("word/document.xml"))
    assert ruby_root.find("w:body/w:p/w:pPr/w:jc", namespace).get(attribute + "val") == "distribute"
    rubies = ruby_root.findall("w:body/w:p/w:r/w:ruby", namespace)
    assert len(rubies) == 2
    assert [item.find("w:rt/w:r/w:t", namespace).text for item in rubies] == ["zhōng", "wén"]
    assert [item.find("w:rubyBase/w:r/w:t", namespace).text for item in rubies] == ["中", "文"]
    assert all(item.find("w:rubyPr/w:rubyAlign", namespace).get(attribute + "val") == "center" for item in rubies)
print("Phonetic guide independently verified as OOXML ruby annotations, separate from body text.")

cell_path = path.with_name("word-cell-interop.docx")
cell_document = Document(cell_path)
assert len(cell_document.tables) == 1
assert cell_document.tables[0].cell(0, 0).text == "底色、垂直居中和独立内边距"
assert int(cell_document.tables[0].cell(0, 0).vertical_alignment) == 1
with ZipFile(cell_path) as edited, ZipFile(path.with_name("word-cell-source.docx")) as original:
    assert edited.testzip() is None
    assert set(edited.namelist()) == set(original.namelist())
    assert [name for name in original.namelist() if original.read(name) != edited.read(name)] == ["word/document.xml"]
    root = ET.fromstring(edited.read("word/document.xml"))
    rows = root.findall("w:body/w:tbl/w:tr", ns)
    for row in rows:
        properties = row.find("w:tc/w:tcPr", ns)
        assert properties.find("w:shd", ns).get(w + "fill").upper() == "CCDDEE"
        assert properties.find("w:vAlign", ns).get(w + "val") == "center"
        for edge, twips in [("left", 370), ("top", 245), ("right", 195), ("bottom", 480)]:
            margin = properties.find("w:tcMar/w:" + edge, ns)
            assert margin.get(w + "w") == str(twips) and margin.get(w + "type") == "dxa"
    assert rows[0].find("w:tc/w:tcPr/w:vMerge", ns).get(w + "val") == "restart"
    assert rows[1].find("w:tc/w:tcPr/w:vMerge", ns) is not None
print("Cell fill, alignment, four margins, merged continuation and unrelated package parts independently verified.")
