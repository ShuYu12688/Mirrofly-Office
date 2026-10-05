"""Independently evaluate the actual bridge DOCX numbering and PDF text."""

from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import re
import sys

from docx import Document
from pypdf import PdfReader


directory = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[2] / 'build'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W = '{' + NS['w'] + '}'


def value(node, tag, default=None):
    child = node.find('w:' + tag, NS)
    return child.get(W + 'val', default) if child is not None else default


def alpha(number):
    result = ''
    while number:
        number, index = divmod(number - 1, 26)
        result = chr(65 + index) + result
    return result


def roman(number):
    result = ''
    for amount, text in ((1000, 'M'), (900, 'CM'), (500, 'D'), (400, 'CD'), (100, 'C'),
                         (90, 'XC'), (50, 'L'), (40, 'XL'), (10, 'X'), (9, 'IX'),
                         (5, 'V'), (4, 'IV'), (1, 'I')):
        count, number = divmod(number, amount)
        result += text * count
    return result


def read(name):
    path = directory / name
    doc = Document(path)
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        body = ET.fromstring(archive.read('word/document.xml')).find('w:body', NS)
        numbering = ET.fromstring(archive.read('word/numbering.xml'))
    definitions = {node.get(W + 'abstractNumId'): node for node in numbering.findall('w:abstractNum', NS)}
    numbers = {node.get(W + 'numId'): node for node in numbering.findall('w:num', NS)}
    assert len(numbers) == len(numbering.findall('w:num', NS))
    counters, labels = {}, []
    for paragraph in body.findall('w:p', NS):
        props = paragraph.find('w:pPr/w:numPr', NS)
        if props is None or value(props, 'numId') == '0':
            labels.append('')
            continue
        identity, level_index = value(props, 'numId'), value(props, 'ilvl', '0')
        number = numbers[identity]
        abstract = definitions[value(number, 'abstractNumId')]
        level = next(n for n in abstract.findall('w:lvl', NS) if n.get(W + 'ilvl') == level_index)
        override = next((n for n in number.findall('w:lvlOverride', NS)
                         if n.get(W + 'ilvl') == level_index), None)
        if override is not None and override.find('w:lvl', NS) is not None:
            level = override.find('w:lvl', NS)
        start = int(value(level, 'start', '0'))
        if override is not None:
            start = int(value(override, 'startOverride', str(start)))
        key = identity, level_index
        counters[key] = counters[key] + 1 if key in counters else start
        number_format, template = value(level, 'numFmt'), value(level, 'lvlText')
        count = counters[key]
        if number_format == 'bullet':
            label = template
        else:
            label = (alpha(count) if number_format.endswith('Letter') else
                     roman(count) if number_format.endswith('Roman') else str(count))
            if number_format.startswith('lower'):
                label = label.lower()
            label = template.replace('%' + str(int(level_index) + 1), label)
        labels.append(label)
    return doc, body, numbering, labels


doc, _, _, labels = read('word-list-styles.docx')
assert [p.text for p in doc.paragraphs] == ['first', 'second', 'keep']
assert labels == ['h.', 'i.', '']
pdf = PdfReader(directory / 'word-list-styles.pdf')
text = ' '.join(page.extract_text() for page in pdf.pages)
assert re.sub(r'\s+', ' ', text).strip() == 'h. first i. second keep'
original, before, old_numbers, old_labels = read('word-list-original.docx')
changed, after, new_numbers, new_labels = read('word-list-imported.docx')
assert [p.text for p in original.paragraphs] == [p.text for p in changed.paragraphs]
assert old_labels == ['(V)', '', '(VI)', 'AA.']
assert new_labels == ['(VIII)', '', '(V)', 'AA.']


def semantic(node):
    return node.tag, sorted(node.attrib.items()), (node.text or '').strip(), [semantic(n) for n in node]


for node in old_numbers:
    key = W + ('numId' if node.tag == W + 'num' else 'abstractNumId')
    preserved = next(n for n in new_numbers if n.tag == node.tag and n.get(key) == node.get(key))
    assert semantic(node) == semantic(preserved)
for old, new in zip(list(before)[1:], list(after)[1:]):
    assert semantic(old) == semantic(new)
assert after[0].get('{urn:keep}keep') == 'paragraph'
assert after[0].find('w:pPr/w:keepNext', NS) is not None
assert changed.paragraphs[0].runs[0].bold
print('Word lists: independent DOCX counters, abstract/override preservation, native PDF labels, parent/AI output passed.')
