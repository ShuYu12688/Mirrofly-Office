"""Generate a small independent DOCX fixture; python-docx is a developer-only dependency."""

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


document = Document()
document.add_heading("独立 DOCX 样例", level=1)
paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = paragraph.add_run("中英混排 — Office 2026")
run.font.name = "Arial"
run.font.size = Pt(16)
run.bold = True
run.italic = True
run.underline = True
fonts = run._element.get_or_add_rPr().rFonts
fonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "宋体")
paragraph.add_run().add_break()
paragraph.add_run("第二行保留空格  与制表\t内容")
table = document.add_table(rows=2, cols=2)
for cell, text in zip([cell for row in table.rows for cell in row.cells], ["事项", "进度", "验证", "完成"]):
    cell.text = text
document.sections[0].header.paragraphs[0].text = "页眉不属于基础正文预览"
document.add_paragraph("尾段")
document.save(Path(__file__).with_name("word-python-docx.docx"))
